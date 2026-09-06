"""Resume text extraction for PDF and DOCX files."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

import fitz  # PyMuPDF
from docx import Document


class ResumeParseError(Exception):
    """Raised when a resume cannot be parsed."""


SUPPORTED_EXTENSIONS = {".pdf", ".docx"}


def _clean_text(text: str) -> str:
    """Normalize extracted text without changing its meaning."""
    lines = [line.strip() for line in text.splitlines()]
    non_empty_lines = [line for line in lines if line]
    return "\n".join(non_empty_lines).strip()


def extract_pdf_text(file_bytes: bytes) -> str:
    """Extract text from all pages of a PDF."""
    try:
        document = fitz.open(stream=file_bytes, filetype="pdf")
    except Exception as exc:
        raise ResumeParseError(
            "The PDF could not be opened. Make sure it is a valid, readable PDF."
        ) from exc

    try:
        text = "\n".join(page.get_text("text") for page in document)
    finally:
        document.close()

    cleaned = _clean_text(text)
    if not cleaned:
        raise ResumeParseError(
            "The PDF contains no readable text. A scanned/image-only PDF may need OCR before analysis."
        )
    return cleaned


def extract_docx_text(file_bytes: bytes) -> str:
    """Extract text from paragraphs and table cells in a DOCX document."""
    try:
        document = Document(BytesIO(file_bytes))
    except Exception as exc:
        raise ResumeParseError(
            "The DOCX could not be opened. Make sure it is a valid Word document."
        ) from exc

    chunks: list[str] = []

    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            chunks.append(paragraph.text)

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            row_text = " | ".join(cell for cell in cells if cell)
            if row_text:
                chunks.append(row_text)

    cleaned = _clean_text("\n".join(chunks))
    if not cleaned:
        raise ResumeParseError(
            "The DOCX contains no readable text."
        )
    return cleaned


def extract_resume_text(file_bytes: bytes, filename: str) -> str:
    """Validate the extension and extract text from a supported resume."""
    if not file_bytes:
        raise ResumeParseError("The uploaded file is empty.")

    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise ResumeParseError("Unsupported file format. Please upload a PDF or DOCX resume.")

    if suffix == ".pdf":
        return extract_pdf_text(file_bytes)

    return extract_docx_text(file_bytes)
